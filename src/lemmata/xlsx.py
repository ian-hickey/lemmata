"""Translate module formulas to and from the .xlsx file format, and inject them.

In the file, Excel stores a LAMBDA as `_xlfn.LAMBDA(_xlpm.x, ...)`: every
function introduced after Excel 2007 carries `_xlfn.` and every LAMBDA
parameter or LET variable carries `_xlpm.`. Without the prefixes the name
opens as #NAME?.
"""

from __future__ import annotations

import re
from pathlib import Path

from .canonical import IDENT_RE, mask_strings, split_args, strip_leading_equals
from .future_functions import FUTURE_FUNCTIONS
from .module import Module

CALL_RE = re.compile(r"(?i)(?<![A-Za-z0-9_.])(?:_xlfn\.)?(LAMBDA|LET)\(")
TOKEN_RE = re.compile(r"(?<![A-Za-z0-9_.#$])([A-Za-z_][A-Za-z0-9_.]*)(\()?")
PREFIX_RE = re.compile(r"_xlfn\.|_xlpm\.|_xlws\.")
LABEL_RE = re.compile(r"^\s*(?P<id>[a-z][a-z0-9_]*)@(?P<version>\d+\.\d+\.\d+) sha256:(?P<hash>[0-9a-f]{64})\s*$")


def declared_names(formula: str) -> set[str]:
    """Lower-cased names declared as LAMBDA parameters or LET variables, at any depth."""
    s = strip_leading_equals(formula)
    masked = mask_strings(s)
    names: set[str] = set()
    for m in CALL_RE.finditer(masked):
        spans = split_args(masked, m.end() - 1)
        if m.group(1).upper() == "LAMBDA":
            candidates = spans[:-1]
        else:
            candidates = spans[0:-1:2]
        for a, b in candidates:
            tok = s[a:b].strip()
            if tok.lower().startswith("_xlpm."):
                tok = tok[6:]
            if IDENT_RE.fullmatch(tok):
                names.add(tok.lower())
    return names


def to_file_formula(formula: str) -> str:
    """Add the _xlfn. and _xlpm. prefixes the file format requires."""
    s = strip_leading_equals(formula)
    names = declared_names(s)
    masked = mask_strings(s)
    out: list[str] = []
    last = 0
    for m in TOKEN_RE.finditer(masked):
        ident, is_call = m.group(1), m.group(2)
        replacement = None
        if is_call:
            stored = FUTURE_FUNCTIONS.get(ident.upper())
            if stored:
                replacement = stored + "("
        elif ident.lower() in names:
            replacement = "_xlpm." + ident
        if replacement is not None:
            out.append(s[last:m.start()])
            out.append(replacement)
            last = m.end()
    out.append(s[last:])
    return "".join(out)


def from_file_formula(text: str) -> str:
    """Remove the file-format prefixes outside string literals."""
    s = strip_leading_equals(text)
    masked = mask_strings(s)
    out: list[str] = []
    last = 0
    for m in PREFIX_RE.finditer(masked):
        out.append(s[last:m.start()])
        last = m.end()
    out.append(s[last:])
    return "".join(out)


def parse_label(comment: str | None) -> dict | None:
    if not comment:
        return None
    m = LABEL_RE.match(comment)
    return m.groupdict() if m else None


def read_defined_names(path: Path | str) -> dict[str, tuple[str, str | None]]:
    """Workbook-scoped defined names as {name: (stored formula, comment)}."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    return {name: (dn.attr_text or "", dn.comment) for name, dn in wb.defined_names.items()}


def inject(path: Path | str, modules: list[Module], out: Path | str | None = None) -> list[str]:
    """Write each module (dependencies already resolved) into the workbook as a defined name.

    Creates the workbook when the path does not exist. Returns the names written.
    Note: openpyxl rewrites the file and drops features it does not model, such
    as charts and macros. Keep a copy of anything you care about.
    """
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    path = Path(path)
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
    written = []
    for mod in modules:
        wb.defined_names[mod.name] = DefinedName(
            mod.name, attr_text=to_file_formula(mod.formula), comment=mod.label
        )
        written.append(mod.name)
    wb.save(out or path)
    return written
