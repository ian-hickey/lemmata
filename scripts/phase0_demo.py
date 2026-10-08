"""Build examples/phase0.xlsx with all three modules, then evaluate and verify it.

Open the file in Excel afterwards: the Phase 0 exit test is that B2:B4 show the
same numbers IronCalc prints here and that `lemma verify` still passes after Excel
saves the file.
"""

from pathlib import Path

import ironcalc
import openpyxl

from lemmata.registry import Registry
from lemmata.verify import format_report, verify_workbook
from lemmata.xlsx import inject

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "examples" / "phase0.xlsx"


def main() -> None:
    reg = Registry.default()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demo"
    ws.append(["Module", "Result", "Inputs"])
    ws.append(["LEMMA.CAGR(100, 200, 10)", "=LEMMA.CAGR(E2, F2, G2)", None, None, 100, 200, 10])
    ws.append(["LEMMA.LOAN_PAYMENT(200000, 5%/12, 360)", "=LEMMA.LOAN_PAYMENT(E3, F3/12, G3)", None, None, 200000, 0.05, 360])
    ws.append(["LEMMA.NPV(10%, E5:E10, first_period 0)", "=LEMMA.NPV(F5, E5:E10, G5)", None, None, -40000, 0.08, 0])
    for v in [8000, 9200, 10000, 12000, 14500]:
        ws.append([None, None, None, None, v])
    ws["B4"] = "=LEMMA.NPV(F4, E4:E9, G4)"
    ws["A4"] = "LEMMA.NPV(8%, E4:E9, first_period 0)"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)

    mods = reg.with_dependencies([reg.get("cagr"), reg.get("loan_payment"), reg.get("npv")])
    inject(OUT, mods)
    print(f"wrote {OUT.relative_to(ROOT)} with {', '.join(m.name for m in mods)}")

    m = ironcalc.load_from_xlsx(str(OUT), "en", "UTC")
    m.evaluate()
    for row in (2, 3, 4):
        print(f"  IronCalc {m.get_cell_value(0, row, 1):40s} = {m.get_cell_value(0, row, 2)}")
    print(format_report(verify_workbook(OUT, reg)))


if __name__ == "__main__":
    main()
