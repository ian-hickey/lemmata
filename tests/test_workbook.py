"""Inject modules into a workbook, read it back through IronCalc, and verify it."""

import openpyxl
import ironcalc
import pytest
from openpyxl.workbook.defined_name import DefinedName

from lemmata.registry import Registry
from lemmata.verify import verify_workbook
from lemmata.xlsx import inject

REG = Registry.default()


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "model.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["A2"], ws["A3"] = 100, 200, 10
    ws["B1"] = "=LEMMA.CAGR(A1, A2, A3)"
    ws["B2"] = "=LEMMA.LOAN_PAYMENT(200000, 0.05/12, 360)"
    for i, v in enumerate([-100, 30, 40, 50], start=1):
        ws.cell(row=i, column=4, value=v)
    ws["B3"] = "=LEMMA.NPV(0.1, D1:D4, 0)"
    wb.save(path)
    inject(path, REG.with_dependencies([REG.get("cagr"), REG.get("loan_payment"), REG.get("npv")]))
    return path


def test_ironcalc_evaluates_injected_modules(workbook):
    m = ironcalc.load_from_xlsx(str(workbook), "en", "UTC")
    m.evaluate()
    assert m.get_cell_value(0, 1, 2) == pytest.approx(0.07177346253629313)
    assert m.get_cell_value(0, 2, 2) == pytest.approx(1073.6432460242797)
    assert m.get_cell_value(0, 3, 2) == pytest.approx(-2.1036814425244366)


def test_verify_reports_current(workbook):
    report = verify_workbook(workbook, REG)
    assert report["ok"]
    assert {e["name"] for e in report["names"]} == {"LEMMA.CAGR", "LEMMA.LOAN_PAYMENT", "LEMMA.NPV"}
    assert all(e["status"] == "current" for e in report["names"])


def test_verify_flags_hand_edited_formula(workbook):
    wb = openpyxl.load_workbook(workbook)
    dn = wb.defined_names["LEMMA.CAGR"]
    wb.defined_names["LEMMA.CAGR"] = DefinedName("LEMMA.CAGR", attr_text=dn.attr_text.replace("- 1", "- 2"), comment=dn.comment)
    wb.save(workbook)
    report = verify_workbook(workbook, REG)
    assert not report["ok"]
    assert {e["name"]: e["status"] for e in report["names"]}["LEMMA.CAGR"] == "modified"


def test_verify_flags_missing_and_wrong_labels(workbook):
    wb = openpyxl.load_workbook(workbook)
    npv = wb.defined_names["LEMMA.NPV"]
    wb.defined_names["LEMMA.NPV"] = DefinedName("LEMMA.NPV", attr_text=npv.attr_text, comment=None)
    cagr = wb.defined_names["LEMMA.CAGR"]
    wb.defined_names["LEMMA.CAGR"] = DefinedName("LEMMA.CAGR", attr_text=cagr.attr_text, comment=REG.get("npv").label)
    wb.save(workbook)
    statuses = {e["name"]: e["status"] for e in verify_workbook(workbook, REG)["names"]}
    assert statuses["LEMMA.NPV"] == "unlabeled"
    assert statuses["LEMMA.CAGR"] == "label_mismatch"
    assert statuses["LEMMA.LOAN_PAYMENT"] == "current"


def test_verify_flags_advisory(workbook, tmp_path, monkeypatch):
    reg = Registry(REG.root)
    mod = reg.get("npv")
    advisories_dir = tmp_path / "advisories"
    advisories_dir.mkdir()
    (advisories_dir / "2026-001.yaml").write_text(
        f"module_id: npv\nmodule_hash: {mod.module_hash}\nseverity: high\nsummary: test advisory\nfixed_version: null\n"
    )
    reg.advisories_dir = advisories_dir
    statuses = {e["name"]: e["status"] for e in verify_workbook(workbook, reg)["names"]}
    assert statuses["LEMMA.NPV"] == "advisory"


def test_verify_reports_outdated_and_advisory_from_history(workbook, tmp_path):
    """A superseded formula is recognised from the published-hash history, and an advised one is flagged."""
    import json

    from lemmata.canonical import formula_hash

    reg = Registry(REG.root)
    cagr = reg.get("cagr")
    old_formula = cagr.formula.replace("- 1", "- 1 + 0")  # a different canonical form, as an old version would have
    old_fh = formula_hash(old_formula)
    old_mh = "ab" * 32
    (tmp_path / "history.json").write_text(json.dumps({"published": [
        {"id": "cagr", "name": "LEMMA.CAGR", "version": "0.0.9", "formula_hash": old_fh, "module_hash": old_mh, "published": "2026-01-01T00:00:00+00:00"},
    ]}))
    reg.root = tmp_path  # history lives beside the registry root
    reg.modules_dir = REG.modules_dir
    reg.advisories_dir = tmp_path / "advisories"
    reg.advisories_dir.mkdir()

    wb = openpyxl.load_workbook(workbook)
    dn = wb.defined_names["LEMMA.CAGR"]
    from lemmata.xlsx import to_file_formula

    wb.defined_names["LEMMA.CAGR"] = DefinedName("LEMMA.CAGR", attr_text=to_file_formula(old_formula), comment=f"cagr@0.0.9 sha256:{old_mh}")
    wb.save(workbook)
    statuses = {e["name"]: e for e in verify_workbook(workbook, reg)["names"]}
    assert statuses["LEMMA.CAGR"]["status"] == "outdated"
    assert "current version is 0.1.0" in statuses["LEMMA.CAGR"]["message"]

    (reg.advisories_dir / "2026-009.yaml").write_text(f"module_id: cagr\nmodule_hash: {old_mh}\nseverity: high\nsummary: old cagr is wrong\nfixed_version: 0.1.0\n")
    reg2 = Registry(REG.root)
    reg2.root, reg2.modules_dir, reg2.advisories_dir = tmp_path, REG.modules_dir, reg.advisories_dir
    report = verify_workbook(workbook, reg2)
    assert {e["name"]: e["status"] for e in report["names"]}["LEMMA.CAGR"] == "advisory"
    assert not report["ok"]
