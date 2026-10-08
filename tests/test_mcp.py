"""Phase 2 exit test: a client with only the MCP server builds a workbook from three modules and verify passes."""

import json
import sys

import anyio
import openpyxl
import pytest
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

from lemmata.registry import Registry

REG = Registry.default()


def _payload(result):
    if result.structured_content:
        sc = result.structured_content
        return sc["result"] if isinstance(sc, dict) and set(sc) == {"result"} else sc
    return json.loads(result.content[0].text)


async def _agent_session(tmp_path):
    params = StdioServerParameters(command=sys.executable, args=["-m", "lemmata.mcp_server"])
    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as session:
            await session.initialize()
            tools = {t.name for t in (await session.list_tools()).tools}
            assert tools == {"search_modules", "get_module", "insert_modules", "verify_workbook"}

            hits = _payload(await session.call_tool("search_modules", {"query": "compound annual growth"}))
            assert hits[0]["id"] == "cagr"
            assert "formula" not in hits[0]  # search returns fixed summary fields only

            detail = _payload(await session.call_tool("get_module", {"module": "npv"}))
            assert detail["name"] == "LEMMA.NPV" and detail["formula_xlsx"].startswith("_xlfn.LAMBDA(")
            assert detail["examples"][0]["inputs"][1] == {"range": [-100, 30, 40, 50]}

            # The agent writes its own cells with any library, then inserts the modules.
            path = tmp_path / "agent.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"], ws["A2"], ws["A3"] = 100, 200, 10
            ws["B1"] = "=LEMMA.CAGR(A1, A2, A3)"
            ws["B2"] = "=LEMMA.LOAN_PAYMENT(200000, 0.05/12, 360)"
            for i, v in enumerate([-100, 30, 40, 50], start=1):
                ws.cell(row=i, column=4, value=v)
            ws["B3"] = "=LEMMA.NPV(0.1, D1:D4, 0)"
            ws["B4"] = "=LEMMA.MISSING(1)"
            wb.save(path)

            inserted = _payload(await session.call_tool("insert_modules", {"workbook": str(path), "modules": ["cagr", "loan_payment", "npv"]}))
            assert inserted["written"] == ["LEMMA.CAGR", "LEMMA.LOAN_PAYMENT", "LEMMA.NPV"]

            report = _payload(await session.call_tool("verify_workbook", {"workbook": str(path)}))
            assert report["undefined"] == ["LEMMA.MISSING"]
            assert not report["ok"]

            # The agent acts on the report: remove the bad cell and verify again.
            wb = openpyxl.load_workbook(path)
            wb.active["B4"] = None
            wb.save(path)
            report = _payload(await session.call_tool("verify_workbook", {"workbook": str(path)}))
            assert report["ok"], report
            assert [e["status"] for e in report["names"]] == ["current", "current", "current"]
            uses = {u["cell"]: u["arguments"] for u in report["uses"]}
            assert uses["B3"] == ["0.1", "D1:D4", "0"]
            return path


def test_agent_builds_workbook_through_mcp_only(tmp_path):
    path = anyio.run(_agent_session, tmp_path)
    import ironcalc

    m = ironcalc.load_from_xlsx(str(path), "en", "UTC")
    m.evaluate()
    assert m.get_cell_value(0, 1, 2) == pytest.approx(0.07177346253629313)
    assert m.get_cell_value(0, 3, 2) == pytest.approx(-2.1036814425244366)
